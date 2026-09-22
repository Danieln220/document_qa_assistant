"""
Reading documents.

One function, `load_document()`, turns any supported file into a list of `Segment`
objects. A segment is one page of a PDF, or one heading section of a DOCX or TXT
file, because those formats have no fixed pages.

Everything downstream (chunking, search, citations) works with segments, so adding
a new file format later means adding one loader function here and nothing else.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import pymupdf
import pytesseract
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from PIL import Image

# A PDF page holding less text than this is treated as a scanned image and sent
# to OCR. A genuine text page always has far more; a scanned page usually has 0.
MIN_CHARS_FOR_TEXT_PAGE = 120

# Resolution used when turning a scanned page into an image for OCR. 300 dpi is
# the usual sweet spot: below it small print breaks up, above it is slower for
# no real gain.
OCR_DPI = 300

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}

# Lines that appear on every page of a generated document and carry no content.
# They are dropped so they never end up inside a quoted citation.
BOILERPLATE_PATTERNS = [
    re.compile(r"^Page \d+ of \d+$", re.I),
    re.compile(r"^.{0,80}\|\s*Uncontrolled when printed$", re.I),
    # The document code, anywhere in the line. OCR often merges the header's
    # left and right halves into a single line, so an anchored pattern misses it.
    re.compile(r"RER-[A-Z]{3}-\d{3}\s+Rev\.?\s*\d+", re.I),
]

# A line this long was almost certainly wrapped at the page margin, so the next
# line continues the same sentence. Table cells and headings are shorter, and
# stay on their own line.
WRAPPED_LINE_MIN_CHARS = 55

# A heading such as "4. Cleaning", "5.1 Vacation" or "2.10 Do I have to refuel it?".
# Used to label chunks that come from the middle of a long page or section.
HEADING_RE = re.compile(r"^\d+(\.\d+)*\.?\s+\S.{0,70}$")


@dataclass
class Segment:
    """One page (PDF) or one heading section (DOCX / TXT) of a document."""

    text: str
    page: int | None          # 1-based page number, or None when the format has no pages
    heading: str | None       # nearest heading, used in citations and to help search
    needed_ocr: bool = False  # True when the text came from OCR rather than a text layer


def load_document(path: Path) -> list[Segment]:
    """Read any supported file into segments. Unsupported files raise ValueError."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    if suffix in (".txt", ".md"):
        return _load_text(path)
    raise ValueError(f"Unsupported file type: {path.name}")


# ---------------------------------------------------------------------------
# PDF (with OCR fallback for scanned pages)
# ---------------------------------------------------------------------------

def _load_pdf(path: Path) -> list[Segment]:
    """
    One segment per page. Pages with a text layer are read directly; pages that
    are just an image (a scan or a photo) are rendered and read by OCR, so the
    page number on the citation is correct either way.
    """
    segments: list[Segment] = []
    heading = None
    with pymupdf.open(path) as doc:
        for page_number, page in enumerate(doc, start=1):
            raw = page.get_text("text")
            needed_ocr = len(raw.strip()) < MIN_CHARS_FOR_TEXT_PAGE
            if needed_ocr:
                raw = _ocr_page(page)

            text, last_heading = _clean_block(raw)
            heading = last_heading or heading  # carry the heading over a page break
            if text:
                segments.append(Segment(text=text, page=page_number, heading=heading, needed_ocr=needed_ocr))
    return segments


def _ocr_page(page: pymupdf.Page) -> str:
    """Render one PDF page to a greyscale image and read it with Tesseract."""
    pix = page.get_pixmap(dpi=OCR_DPI, colorspace=pymupdf.csGRAY)
    image = Image.frombytes("L", (pix.width, pix.height), pix.samples)
    return pytesseract.image_to_string(image)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def _load_docx(path: Path) -> list[Segment]:
    """
    One segment per heading section. Word files have no fixed pages (the page
    breaks depend on the reader's printer and zoom), so citations for DOCX show
    the section heading instead of a page number.
    """
    document = Document(path)
    segments: list[Segment] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            segments.append(Segment(text="\n".join(buffer).strip(), page=None, heading=heading))
            buffer.clear()

    for block in _iter_docx_blocks(document):
        if isinstance(block, DocxTable):
            # A table belongs to the section it appears in, so it is written into
            # the current section's text, one row per line. Reading tables in
            # document order is what keeps, for example, the vacation table under
            # the "Vacation" heading rather than at the end of the file.
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in block.rows]
            buffer.extend(rows)
            continue

        text = block.text.strip()
        if not text:
            continue
        if block.style.name.startswith("Heading") or block.style.name == "Title":
            flush()                      # the previous section ends where the next heading starts
            heading = text
            buffer.append(text)          # keep the heading in the text so it can be searched
        else:
            buffer.append(text)
    flush()
    return segments


def _iter_docx_blocks(document: Document):
    """
    Yield the document's paragraphs and tables in the order they appear.

    python-docx exposes `.paragraphs` and `.tables` as two separate lists, which
    loses their relative order; walking the underlying XML body keeps it.
    """
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, document)


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------

def _load_text(path: Path) -> list[Segment]:
    """
    One segment per heading section. Headings are either underlined with '=' or
    '-' (as our generated text files are) or numbered, like "2.5 Do you deliver?".
    """
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    segments: list[Segment] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        # Text files are wrapped at a fixed width, so they need the same
        # line-joining treatment as a PDF.
        body, _ = _clean_block("\n".join(buffer))
        if body.strip():
            segments.append(Segment(text=body, page=None, heading=heading))
        buffer.clear()

    for i, line in enumerate(lines):
        stripped = line.strip()
        underlined = (
            stripped
            and i + 1 < len(lines)
            and set(lines[i + 1].strip()) in ({"="}, {"-"})
            and len(lines[i + 1].strip()) >= len(stripped) - 1
        )
        if underlined or (stripped and HEADING_RE.match(stripped)):
            flush()
            heading = stripped
            buffer.append(stripped)
        elif set(stripped) in ({"="}, {"-"}) and stripped:
            continue  # the underline itself
        else:
            buffer.append(line)
    flush()
    return segments


# ---------------------------------------------------------------------------
# Shared cleanup
# ---------------------------------------------------------------------------

def _clean_block(raw: str) -> tuple[str, str | None]:
    """
    Drop repeated headers and footers, join lines that were wrapped at the page
    margin, and return the text with the last heading seen in it.

    Joining matters for citations: a sentence broken across two lines cannot be
    quoted or verified as one phrase later on.
    """
    kept: list[str] = []
    last_heading: str | None = None
    for line in raw.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if any(p.search(stripped) for p in BOILERPLATE_PATTERNS):
            continue
        if HEADING_RE.match(stripped):
            last_heading = stripped
            kept.append(stripped)
            continue
        # Continue the previous line when that line looks wrapped rather than finished.
        if (
            kept
            and len(kept[-1]) >= WRAPPED_LINE_MIN_CHARS
            and not kept[-1].endswith((".", ":", ";", "!", "?"))
            and not HEADING_RE.match(kept[-1])
        ):
            kept[-1] = f"{kept[-1]} {stripped}"
        else:
            kept.append(stripped)
    return "\n".join(kept), last_heading


def find_heading(text: str, first_lines: int | None = None) -> str | None:
    """
    The last heading line inside a piece of text, if any.

    With `first_lines`, only that many lines from the start are examined, which
    answers a different question: "does this chunk open with a heading?"
    """
    lines = text.splitlines()
    if first_lines is not None:
        lines = lines[:first_lines]
    found = None
    for line in lines:
        stripped = line.strip()
        if HEADING_RE.match(stripped):
            found = stripped
    return found
